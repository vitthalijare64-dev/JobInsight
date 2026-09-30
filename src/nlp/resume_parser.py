from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Dict, List

import re

import fitz
from docx import Document


# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

SKILL_DATA_PATH = (
    PROJECT_ROOT
    / "data"
    / "processed"
    / "job_skills_normalized.parquet"
)


# ============================================================
# FALLBACK SKILLS
# ============================================================

FALLBACK_SKILLS = [
    "Python",
    "Java",
    "JavaScript",
    "TypeScript",
    "C++",
    "C#",
    "SQL",
    "Excel",
    "Power BI",
    "Tableau",
    "AWS",
    "Azure",
    "Google Cloud",
    "Docker",
    "Kubernetes",
    "Git",
    "Machine Learning",
    "Deep Learning",
    "Artificial Intelligence",
    "NLP",
    "Data Analysis",
    "Data Analytics",
    "Data Science",
    "Statistics",
    "Pandas",
    "NumPy",
    "scikit-learn",
    "TensorFlow",
    "PyTorch",
    "Spark",
    "Hadoop",
    "ETL",
    "REST API",
    "API",
    "Agile",
    "Scrum",
    "Project Management",
]


# ============================================================
# RESUME SECTIONS
# ============================================================

SECTION_NAMES = {
    "summary": [
        "summary",
        "professional summary",
        "profile",
        "professional profile",
        "career objective",
        "objective",
    ],
    "experience": [
        "experience",
        "work experience",
        "professional experience",
        "employment history",
        "work history",
    ],
    "education": [
        "education",
        "academic background",
        "academic credentials",
        "qualifications",
    ],
    "skills": [
        "skills",
        "technical skills",
        "core skills",
        "technical competencies",
        "core competencies",
    ],
    "projects": [
        "projects",
        "academic projects",
        "personal projects",
        "key projects",
    ],
    "certifications": [
        "certifications",
        "certificates",
        "professional certifications",
        "licenses",
    ],
}


# ============================================================
# CANDIDATE PROFILE
# ============================================================

@dataclass
class CandidateProfile:
    name: str
    email: str
    phone: str
    experience: str
    education: List[str]
    skills: List[str]
    sections: Dict[str, str]
    certifications: List[str]
    projects: List[str]
    word_count: int
    raw_text: str

    def to_dict(self):
        return asdict(self)


# ============================================================
# SKILL LOADING
# ============================================================

def load_market_skills() -> List[str]:
    """
    Load skills discovered from the JobInsight job-market dataset.

    Falls back to the built-in skill dictionary if the processed
    skill dataset is unavailable.
    """

    skills = set(FALLBACK_SKILLS)

    try:
        import pandas as pd

        if not SKILL_DATA_PATH.exists():
            return sorted(
                skills,
                key=len,
                reverse=True,
            )

        df = pd.read_parquet(
            SKILL_DATA_PATH
        )

        if "skills" not in df.columns:
            return sorted(
                skills,
                key=len,
                reverse=True,
            )

        for value in df["skills"].dropna():

            if isinstance(
                value,
                (list, tuple, set),
            ):

                for skill in value:

                    skill = str(
                        skill
                    ).strip()

                    if skill:
                        skills.add(skill)

            else:

                text = str(value)

                for skill in text.split(","):

                    skill = skill.strip()

                    if skill:
                        skills.add(skill)

    except Exception:
        # Parser should remain usable even if the market
        # skill dataset cannot be loaded.
        pass

    return sorted(
        skills,
        key=len,
        reverse=True,
    )


# ============================================================
# FILE EXTRACTION
# ============================================================

def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extract text from a PDF using PyMuPDF.
    """

    document = fitz.open(
        stream=file_bytes,
        filetype="pdf",
    )

    pages = []

    try:

        for page in document:

            text = page.get_text()

            if text:
                pages.append(text)

    finally:

        document.close()

    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX file.
    """

    import io

    document = Document(
        io.BytesIO(file_bytes)
    )

    paragraphs = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            paragraphs.append(text)

    # Also read table content because some resumes
    # place contact information or experience in tables.
    for table in document.tables:

        for row in table.rows:

            row_values = []

            for cell in row.cells:

                cell_text = cell.text.strip()

                if cell_text:
                    row_values.append(
                        cell_text
                    )

            if row_values:
                paragraphs.append(
                    " | ".join(row_values)
                )

    return "\n".join(paragraphs).strip()


def extract_text(
    file_bytes: bytes,
    filename: str,
) -> str:
    """
    Extract text from PDF or DOCX.
    """

    extension = Path(
        filename
    ).suffix.lower()

    if extension == ".pdf":

        return extract_pdf_text(
            file_bytes
        )

    if extension == ".docx":

        return extract_docx_text(
            file_bytes
        )

    raise ValueError(
        "Unsupported resume format. "
        "Only PDF and DOCX are supported."
    )


# ============================================================
# BASIC CONTACT EXTRACTION
# ============================================================

def extract_email(text: str) -> str:

    matches = re.findall(
        r"\b[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text,
    )

    if matches:
        return matches[0]

    return "Not detected"


def extract_phone(text: str) -> str:

    matches = re.findall(
        r"(?:\+?\d[\d\s().-]{8,}\d)",
        text,
    )

    if matches:
        return matches[0].strip()

    return "Not detected"


# ============================================================
# NAME EXTRACTION
# ============================================================

def extract_name(text: str) -> str:

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    if not lines:
        return "Not detected"

    ignored = {
        "resume",
        "curriculum vitae",
        "cv",
        "profile",
        "professional profile",
    }

    for line in lines[:15]:

        lower = line.lower()

        if lower in ignored:
            continue

        if "@" in line:
            continue

        if re.search(
            r"\d",
            line,
        ):
            continue

        words = line.split()

        if not 2 <= len(words) <= 5:
            continue

        valid = all(
            re.match(
                r"^[A-Za-z.'-]+$",
                word,
            )
            for word in words
        )

        if valid:
            return line

    return "Not detected"


# ============================================================
# SECTION DETECTION
# ============================================================

def normalize_heading(text: str) -> str:

    text = text.lower()

    text = re.sub(
        r"[^a-zA-Z ]",
        "",
        text,
    )

    return " ".join(
        text.split()
    ).strip()


def detect_sections(
    text: str,
) -> Dict[str, str]:

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    sections = {
        section: []
        for section in SECTION_NAMES
    }

    current_section = None

    # Reverse lookup of possible headings
    heading_lookup = {}

    for section, names in SECTION_NAMES.items():

        for name in names:

            heading_lookup[
                normalize_heading(name)
            ] = section

    for line in lines:

        normalized = normalize_heading(
            line
        )

        if normalized in heading_lookup:

            current_section = (
                heading_lookup[normalized]
            )

            continue

        if current_section is not None:

            sections[
                current_section
            ].append(line)

    return {
        section: "\n".join(
            values
        ).strip()
        for section, values
        in sections.items()
    }


# ============================================================
# SKILL EXTRACTION
# ============================================================

def extract_skills(
    text: str,
    skills_dictionary: List[str] | None = None,
) -> List[str]:

    if not text:
        return []

    if skills_dictionary is None:
        skills_dictionary = (
            load_market_skills()
        )

    text_lower = text.lower()

    detected = []

    for skill in skills_dictionary:

        skill = str(
            skill
        ).strip()

        if not skill:
            continue

        pattern = (
            r"(?<![a-zA-Z0-9])"
            + re.escape(
                skill.lower()
            )
            + r"(?![a-zA-Z0-9])"
        )

        if re.search(
            pattern,
            text_lower,
        ):

            detected.append(skill)

    # Remove duplicates while preserving order
    detected = list(
        dict.fromkeys(
            detected
        )
    )

    return sorted(
        detected,
        key=str.lower,
    )


# ============================================================
# EDUCATION EXTRACTION
# ============================================================

EDUCATION_PATTERNS = [
    (
        "Doctorate",
        [
            "ph.d",
            "phd",
            "doctorate",
            "doctoral",
            "doctor of",
        ],
    ),
    (
        "Master",
        [
            "master's",
            "masters",
            "master of",
            "m.tech",
            "mtech",
            "m.sc",
            "msc",
            "m.s.",
            "mba",
            "m.e.",
        ],
    ),
    (
        "Bachelor",
        [
            "bachelor's",
            "bachelors",
            "bachelor of",
            "b.tech",
            "btech",
            "b.sc",
            "bsc",
            "b.s.",
            "b.e.",
            "bba",
            "bca",
        ],
    ),
    (
        "Associate",
        [
            "associate degree",
            "associate of",
        ],
    ),
    (
        "High School",
        [
            "high school",
            "secondary school",
            "12th",
            "hsc",
        ],
    ),
]


def detect_education(
    text: str,
    education_section: str = "",
) -> List[str]:

    # Prefer the actual Education section.
    # This prevents certifications elsewhere in the
    # resume from being interpreted as education.
    search_text = (
        education_section
        if education_section.strip()
        else text
    )

    text_lower = search_text.lower()

    detected = []

    for label, patterns in EDUCATION_PATTERNS:

        for pattern in patterns:

            if pattern in text_lower:

                detected.append(
                    label
                )

                break

    return list(
        dict.fromkeys(
            detected
        )
    )


# ============================================================
# EXPERIENCE DETECTION
# ============================================================

def detect_experience(
    text: str,
) -> str:

    patterns = [
        r"(\d+(?:\.\d+)?)\s*\+?\s*years?",
        r"(\d+(?:\.\d+)?)\s*yrs?",
    ]

    values = []

    for pattern in patterns:

        matches = re.findall(
            pattern,
            text.lower(),
        )

        for value in matches:

            try:
                values.append(
                    float(value)
                )
            except ValueError:
                continue

    if not values:
        return "Not detected"

    years = max(values)

    if years < 1:
        return "< 1 year"

    if years < 2:
        return "1+ year"

    return f"{years:g}+ years"


# ============================================================
# PROJECT EXTRACTION
# ============================================================

def extract_projects(
    section_text: str,
) -> List[str]:

    if not section_text.strip():
        return []

    lines = [
        line.strip()
        for line in section_text.splitlines()
        if line.strip()
    ]

    return lines


# ============================================================
# CERTIFICATION EXTRACTION
# ============================================================

def extract_certifications(
    section_text: str,
) -> List[str]:

    if not section_text.strip():
        return []

    lines = [
        line.strip()
        for line in section_text.splitlines()
        if line.strip()
    ]

    return lines


# ============================================================
# MAIN PARSER
# ============================================================

def parse_resume(
    file_bytes: bytes,
    filename: str,
) -> CandidateProfile:
    """
    Parse a PDF or DOCX resume into a structured
    CandidateProfile.
    """

    raw_text = extract_text(
        file_bytes,
        filename,
    )

    if not raw_text.strip():

        raise ValueError(
            "No readable text was detected "
            "in the uploaded resume."
        )

    sections = detect_sections(
        raw_text
    )

    skills = extract_skills(
        raw_text
    )

    education = detect_education(
        raw_text,
        sections.get(
            "education",
            "",
        ),
    )

    certifications = (
        extract_certifications(
            sections.get(
                "certifications",
                "",
            )
        )
    )

    projects = extract_projects(
        sections.get(
            "projects",
            "",
        )
    )

    profile = CandidateProfile(
        name=extract_name(
            raw_text
        ),
        email=extract_email(
            raw_text
        ),
        phone=extract_phone(
            raw_text
        ),
        experience=detect_experience(
            raw_text
        ),
        education=education,
        skills=skills,
        sections=sections,
        certifications=certifications,
        projects=projects,
        word_count=len(
            raw_text.split()
        ),
        raw_text=raw_text,
    )

    return profile


# ============================================================
# SIMPLE TEST
# ============================================================

if __name__ == "__main__":

    print("=" * 70)
    print("JobInsight - Resume Parser")
    print("=" * 70)

    print(
        f"Skill dataset: {SKILL_DATA_PATH}"
    )

    skills = load_market_skills()

    print(
        f"Market skills loaded: {len(skills):,}"
    )

    print(
        "Resume parser ready."
    )

    print("=" * 70)